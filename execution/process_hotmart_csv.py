"""
process_hotmart_csv.py
Processa os exports CSV da Hotmart e salva em .tmp/hotmart_raw.csv

Filtros aplicados:
  - Status in ["Aprovado", "Completo"]
  - Produto: nome contém "plataforma medsimple" (case-insensitive)

Formato dos CSVs:
  - Delimitador: ponto e vírgula (;)
  - Encoding: utf-8-sig
  - Data de Confirmação: DD/MM/YYYY HH:MM:SS
"""

import os
import csv
import shutil
import tempfile
from datetime import datetime

OUTPUT_FILE = os.path.join(os.path.dirname(__file__), "..", ".tmp", "hotmart_raw.csv")

CSV_FILES = [
    r"F:\Downloads\sales_history_20260320154544_FB2456C83882739997474803165.csv",
    r"F:\Downloads\sales_history_20260320154431_473E7C3A5715212713349273767.csv",
    r"F:\Downloads\sales_history_20260320132919_41DAFA167027452998184058871.csv",
    r"F:\Downloads\sales_history_20260320151431_8219BF837955421160596342750.csv",
    r"F:\Downloads\sales_history_20260320132640_8584E6A83289272021253433931.csv",
    r"F:\Downloads\sales_history_20260320154900_5DCEB2B413074293895924279648.csv",
]

XLSX_FILES = [
    r"F:\Downloads\sales_history_20260430174224_89AFB0589533606856776939868.xls",
]

STATUS_VALIDOS = {"aprovado", "completo"}
PRODUTO_FILTRO = "plataforma medsimple"


def normalize_phone(ddd, number):
    digits = "".join(filter(str.isdigit, str(ddd or "") + str(number or "")))
    # Remover DDI +55 se presente
    if digits.startswith("55") and len(digits) >= 12:
        digits = digits[2:]
    return digits


def parse_date(value):
    if not value or str(value).strip() == "":
        return None
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(value).strip(), fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    return None


def find_col(headers, *candidates):
    """
    Encontra o nome original de uma coluna pelos candidatos (case-insensitive).
    Prioriza match exato antes de match parcial.
    """
    h_lower = [h.lower().strip() for h in headers]
    for candidate in candidates:
        cand = candidate.lower()
        for i, h in enumerate(h_lower):
            if h == cand:
                return headers[i]
        for i, h in enumerate(h_lower):
            if cand in h:
                return headers[i]
    return None


def process_file(filepath):
    print(f"\n[Hotmart] Processando: {os.path.basename(filepath)}")

    records = []
    skipped_status = 0
    skipped_produto = 0
    skipped_date = 0

    with open(filepath, encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.DictReader(f, delimiter=";")
        headers = reader.fieldnames or []

        # Encontrar colunas (tolera encoding corrompido nos nomes)
        col_status = find_col(headers, "status")
        col_email = find_col(headers, "email")
        col_nome = find_col(headers, "nome")
        col_ddd = find_col(headers, "ddd")
        col_tel = find_col(headers, "telefone")
        col_data = find_col(headers, "confirma")  # "Data de Confirmação"
        col_produto = find_col(headers, "nome do produto", "produto")
        col_valor = find_col(headers, "faturamento", "valor que você recebeu convertido", "preço da oferta")

        print(f"  status={col_status} | email={col_email} | nome={col_nome} | tel={col_tel} | data={col_data} | produto={col_produto} | valor={col_valor}")

        for row in reader:
            status = str(row.get(col_status, "") or "").strip().lower()
            if status not in STATUS_VALIDOS:
                skipped_status += 1
                continue

            if col_produto:
                produto = str(row.get(col_produto, "") or "").strip().lower()
                if PRODUTO_FILTRO not in produto:
                    skipped_produto += 1
                    continue

            data = parse_date(row.get(col_data, ""))
            if not data:
                skipped_date += 1
                continue

            email = str(row.get(col_email, "") or "").strip().lower()
            nome = str(row.get(col_nome, "") or "").strip()
            ddd = str(row.get(col_ddd, "") or "").strip()
            tel = str(row.get(col_tel, "") or "").strip()
            telefone = normalize_phone(ddd, tel)

            if not email and not telefone:
                continue

            valor_raw = str(row.get(col_valor, "") or "").strip().replace(",", ".") if col_valor else ""
            try:
                valor = float(valor_raw)
            except ValueError:
                valor = 0.0

            nome_produto_raw = str(row.get(col_produto, "") or "").strip() if col_produto else ""

            records.append({
                "email": email,
                "telefone": telefone,
                "nome": nome,
                "data_compra": data,
                "plataforma": "hotmart",
                "valor_liquido": valor,
                "nome_produto": nome_produto_raw,
            })

    print(f"  Válidos: {len(records)} | Ignorados status: {skipped_status} | Produto errado: {skipped_produto} | Sem data: {skipped_date}")
    return records


def process_xlsx_file(filepath):
    """Processa exports XLSX da Hotmart (mesmo schema do CSV, mas em planilha)."""
    import openpyxl
    print(f"\n[Hotmart] Processando XLSX: {os.path.basename(filepath)}")

    # Hotmart às vezes nomeia .xls mas o arquivo é XLSX moderno; copiar pra extensão correta
    tmp_path = filepath
    if not filepath.lower().endswith(".xlsx"):
        tmp_path = os.path.join(tempfile.gettempdir(), os.path.basename(filepath) + ".xlsx")
        shutil.copy(filepath, tmp_path)

    wb = openpyxl.load_workbook(tmp_path, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in ws[1]]

    def col_idx(*candidates):
        h_lower = [h.lower() for h in headers]
        for cand in candidates:
            cl = cand.lower()
            for i, h in enumerate(h_lower):
                if h == cl:
                    return i
            for i, h in enumerate(h_lower):
                if cl in h:
                    return i
        return None

    idx_status = col_idx("status")
    idx_email = col_idx("email")
    idx_nome = col_idx("nome")
    idx_ddd = col_idx("ddd")
    idx_tel = col_idx("telefone")
    idx_data = col_idx("confirma")
    idx_produto = col_idx("nome do produto", "produto")
    idx_valor = col_idx("faturamento", "valor que voc")

    print(f"  status={idx_status} email={idx_email} nome={idx_nome} tel={idx_tel} data={idx_data} produto={idx_produto} valor={idx_valor}")

    records = []
    skipped_status = skipped_produto = skipped_date = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        status = str(row[idx_status] or "").strip().lower() if idx_status is not None else ""
        if status not in STATUS_VALIDOS:
            skipped_status += 1
            continue

        produto = str(row[idx_produto] or "").strip() if idx_produto is not None else ""
        if PRODUTO_FILTRO not in produto.lower():
            skipped_produto += 1
            continue

        data_val = row[idx_data] if idx_data is not None else None
        if isinstance(data_val, datetime):
            data = data_val.strftime("%Y-%m-%d %H:%M:%S")
        else:
            data = parse_date(data_val)
        if not data:
            skipped_date += 1
            continue

        email = str(row[idx_email] or "").strip().lower() if idx_email is not None else ""
        nome = str(row[idx_nome] or "").strip() if idx_nome is not None else ""
        ddd = str(row[idx_ddd] or "").strip() if idx_ddd is not None else ""
        tel = str(row[idx_tel] or "").strip() if idx_tel is not None else ""
        telefone = normalize_phone(ddd, tel)

        if not email and not telefone:
            continue

        valor_raw = row[idx_valor] if idx_valor is not None else 0
        try:
            valor = float(str(valor_raw or "0").replace(",", "."))
        except (ValueError, TypeError):
            valor = 0.0

        records.append({
            "email": email,
            "telefone": telefone,
            "nome": nome,
            "data_compra": data,
            "plataforma": "hotmart",
            "valor_liquido": valor,
            "nome_produto": produto,
        })

    wb.close()
    print(f"  Válidos: {len(records)} | Status: {skipped_status} | Produto: {skipped_produto} | Sem data: {skipped_date}")
    return records


def main():
    print("=" * 55)
    print("HOTMART — Processando CSVs históricos")
    print("=" * 55)

    all_records = []
    for f in CSV_FILES:
        if not os.path.exists(f):
            print(f"[!] Arquivo não encontrado: {f}")
            continue
        all_records.extend(process_file(f))

    for f in XLSX_FILES:
        if not os.path.exists(f):
            print(f"[!] Arquivo não encontrado: {f}")
            continue
        all_records.extend(process_xlsx_file(f))

    # Mesclar com hotmart_raw.csv existente para preservar histórico
    if os.path.exists(OUTPUT_FILE):
        with open(OUTPUT_FILE, "r", encoding="utf-8") as fh:
            reader = csv.DictReader(fh)
            existing = list(reader)
        print(f"\n[Hotmart] Mesclando com {len(existing)} registros existentes")
        for r in existing:
            try:
                r["valor_liquido"] = float(r.get("valor_liquido") or 0)
            except (ValueError, TypeError):
                r["valor_liquido"] = 0.0
            all_records.append(r)

    # Dedupe por (email, telefone, data_compra) — primeira ocorrência ganha (XLSX/CSV novo vem antes)
    seen = set()
    unique = []
    for r in all_records:
        key = (r["email"], r["telefone"], r["data_compra"])
        if key not in seen:
            seen.add(key)
            unique.append(r)

    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    with open(OUTPUT_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["email", "telefone", "nome", "data_compra", "plataforma", "valor_liquido", "nome_produto"])
        writer.writeheader()
        writer.writerows(unique)

    print(f"\n[Hotmart] Total bruto: {len(all_records)} | Únicos: {len(unique)}")
    print(f"[Hotmart] Salvo em: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
